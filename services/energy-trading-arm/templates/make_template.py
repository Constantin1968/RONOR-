"""Generate the daily-operations Excel templates (RO/UA/MD corridors).

Usage:  python templates/make_template.py
Output: templates/daily_ops_template.xlsx        (NTC / prices, market side)
        templates/operatiuni_zilnice_template.xlsx (our position + results, twin side)
"""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

OUT = Path(__file__).resolve().parent / "daily_ops_template.xlsx"
OUT_OPS = Path(__file__).resolve().parent / "operatiuni_zilnice_template.xlsx"
BOLD = Font(bold=True)
CORRIDORS = ["Ua - Md", "Md - Ro", "Ro - Ua", "Ua - Ro", "Transfer"]
GROUPS = ["Capacity won", "CBC Price", "Bid Limit Price", "Nominated MW", "Profit EUR"]


def make_operations_template() -> None:
    """The operator's own sheet: what we won, what we bid, what we executed, what it made.

    Same layout as the auction-result table posted in the group, plus two result
    groups (Nominated MW, Profit EUR) to fill in after delivery. Intervals are CET
    1–24; a corridor column left 0/blank means "not held". Ro/Ua Price are optional —
    the service reads OPCOM/OREE itself and only verifies what is typed here.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Operatiuni"
    head = ["Intervals CET"]
    sub = [""]
    for g in GROUPS:
        head += [g] + [""] * (len(CORRIDORS) - 1)
        sub += CORRIDORS
    head += ["Ro Price", "Ua Price"]
    sub += ["", ""]
    ws.append(head)
    ws.append(sub)
    col = 2
    for _ in GROUPS:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 4)
        ws.cell(row=1, column=col).alignment = Alignment(horizontal="center")
        col += 5
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)
    for c in (*ws[1], *ws[2]):
        c.font = BOLD
    for h in range(1, 25):
        # Auction groups default to 0 (not held); result groups stay blank (= not yet known).
        ws.append([h] + [0] * (3 * len(CORRIDORS)) + [None] * (2 * len(CORRIDORS) + 2))
    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 14

    note = wb.create_sheet("Cum se completeaza")
    for line in (
        "Un fișier pe zi de livrare; pune ziua în numele fișierului sau în mesaj (ex. 'operatiuni 14.09').",
        "Capacity won: MW câștigați pe coridor și interval (0 = nu deținem).",
        "CBC Price: prețul plătit pe capacitate (EUR/MWh), pe același interval.",
        "Bid Limit Price: limita pusă pe piața de energie (RO: preț maxim de cumpărare; UA: preț minim de vânzare).",
        "Nominated MW: cât s-a executat/nominalizat efectiv (după livrare). Gol = încă necunoscut, 0 = nimic.",
        "Profit EUR: rezultatul pe care îl contabilizezi tu pe interval (după livrare) — twin-ul îl compară cu calculul lui.",
        "Ro Price / Ua Price: opțional; serviciul citește singur OPCOM și OREE.",
        "Poți trimite dimineața doar primele 3 grupuri și seara același fișier completat: rândurile se îmbină.",
    ):
        note.append([line])
    note.column_dimensions["A"].width = 120
    wb.save(OUT_OPS)
    print(f"template scris: {OUT_OPS} ({OUT_OPS.stat().st_size} octeți)")


def main() -> None:
    wb = Workbook()

    atc = wb.active
    atc.title = "ATC"
    atc.append(["Granita", "ATC_MW"])
    for row in (["RO-UA", 450], ["UA-MD", 600], ["RO-MD", 400], ["UA-MD-RO", 400]):
        atc.append(row)
    for cell in atc[1]:
        cell.font = BOLD

    preturi = wb.create_sheet("Preturi")
    preturi.append(["Zona", "Ora", "Pret_EUR_MWh"])
    for row in (["RO", 18, 112.5], ["MD", 19, 121.0], ["UA", 18, 68.0]):
        preturi.append(row)
    for cell in preturi[1]:
        cell.font = BOLD

    matrice = wb.create_sheet("Matrice_ore")
    matrice.append(["Zona", *list(range(24))])
    matrice.append(["RO", *([97.0] * 24)])
    matrice.append(["UA", *([70.0] * 24)])
    matrice.append(["MD", *([104.0] * 24)])
    for cell in matrice[1]:
        cell.font = BOLD

    for sheet in (atc, preturi, matrice):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(OUT)
    print(f"template scris: {OUT} ({OUT.stat().st_size} octeți)")
    make_operations_template()


if __name__ == "__main__":
    main()
