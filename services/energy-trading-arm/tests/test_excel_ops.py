from io import BytesIO

from openpyxl import Workbook

from energy_trading.ops_intake import read_excel_ops


def _workbook() -> bytes:
    wb = Workbook()
    atc = wb.active
    atc.title = "ATC"
    atc.append(["Granita", "ATC_MW"])
    atc.append(["RO-UA", 450])
    atc.append(["UA/MD", 600])

    preturi = wb.create_sheet("Preturi")
    preturi.append(["Zona", "Ora", "Pret_EUR_MWh"])
    preturi.append(["RO", 18, 112.5])
    preturi.append(["MD", 19, 121])

    matrice = wb.create_sheet("Matrice_ore")
    matrice.append(["Zona", *list(range(24))])
    matrice.append(["UA", *([68.0] * 24)])

    liber = wb.create_sheet("Liber")
    liber.append(["RO-MD ATC 400 MW"])

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_read_excel_all_layouts():
    intake = read_excel_ops(_workbook(), day="2026-09-12")
    assert intake.availability == {"RO-UA": 450.0, "UA-MD": 600.0, "RO-MD": 400.0}
    assert intake.prices_override["RO"][18] == 112.5
    assert intake.prices_override["MD"][19] == 121.0
    assert intake.prices_override["UA"][7] == 68.0
    assert intake.warnings == []


def test_read_excel_invalid_bytes_warns():
    intake = read_excel_ops(b"not a workbook")
    assert intake.availability == {}
    assert any("invalid" in w for w in intake.warnings)


def test_template_matches_parser():
    from pathlib import Path

    template = Path(__file__).resolve().parent.parent / "templates" / "daily_ops_template.xlsx"
    intake = read_excel_ops(template.read_bytes())
    assert intake.availability["RO-UA"] == 450.0
    assert intake.prices_override["RO"][18] == 112.5
    assert intake.prices_override["UA"][0] == 70.0


def test_operations_template_matches_parser():
    from pathlib import Path

    template = (
        Path(__file__).resolve().parent.parent / "templates" / "operatiuni_zilnice_template.xlsx"
    )
    intake = read_excel_ops(template.read_bytes(), day="2026-09-15")
    assert intake.warnings == []  # the "Cum se completeaza" sheet is skipped
    assert set(intake.bids) == {"capacity", "cbc", "limits"}  # result columns blank = unknown
    assert set(intake.bids["capacity"]) == {"UA-MD", "MD-RO", "RO-UA", "UA-RO", "UA-MD-RO"}
    assert intake.position_summary() == []  # nothing held in the empty template


def test_api_ops_upload():
    from fastapi.testclient import TestClient

    from energy_trading.api import app

    c = TestClient(app)
    r = c.post(
        "/api/ops-upload",
        params={"day": "2026-09-12"},
        files={
            "file": (
                " zilnic.xlsx",
                _workbook(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert r.status_code == 200
    body = r.json()
    assert body["availability"]["UA-MD"] == 600.0
    assert body["prices_override"]["MD"]["19"] == 121.0


def test_api_ops_upload_rejects_non_excel():
    from fastapi.testclient import TestClient

    from energy_trading.api import app

    c = TestClient(app)
    r = c.post("/api/ops-upload", files={"file": ("note.txt", b"hello", "text/plain")})
    assert r.status_code == 400
