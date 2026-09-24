"""The operator's daily sheet — capacity won, CBC, limits, then fills and results —
lands as the twin's ``bids_<day>.csv`` and re-triggers the P/L Digital Twin."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook

from energy_trading.ops_intake import (
    corridor_key,
    load_bids_csv,
    read_csv_ops,
    read_excel_ops,
    write_bids_csv,
)
from energy_trading.twin_pnl import day_pnl, format_pnl

CORRIDORS = ["Ua - Md", "Md - Ro", "Ro - Ua", "Ua - Ro", "Transfer"]


def _operator_sheet(with_results: bool = False) -> bytes:
    """Mirror of the sheet the operator posts: merged group header + corridor row."""
    wb = Workbook()
    ws = wb.active
    ws.title = "14.09"
    groups = ["Capacity won", "CBC Price", "Bid Limit Price"]
    if with_results:
        groups += ["Nominated MW", "Profit EUR"]
    head = ["Intervals CET"]
    sub = [""]
    for g in groups:
        head += [g] + [""] * (len(CORRIDORS) - 1)
        sub += CORRIDORS
    head += ["Ro Price", "Ua Price"]
    sub += ["", ""]
    ws.append(head)
    ws.append(sub)
    col = 2
    for g in groups:
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 4)
        col += 5
    for h in range(1, 25):
        cap = [0, 0, 15, 0, 10]  # RO-UA 15 MW and UA-MD-RO 10 MW every hour
        cbc = [0, 0, 0.23 if h < 5 else 0.0, 0, 1.5]
        lim = [0, 0, 55 + h, 0, 200]
        row = [h, *cap, *cbc, *lim]
        if with_results:
            nominated = [0, 0, 15 if h in (1, 2, 3) else 0, 0, 10]
            profit = [0, 0, 120.5 if h in (1, 2, 3) else 0, 0, -3.0]
            row += [*nominated, *profit]
        row += [100 + h, 95.5]
        ws.append(row)
    ws.append(["Total", "", "", 360])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_corridor_key_normalises_operator_labels():
    assert corridor_key("Ua - Md") == "UA-MD"
    assert corridor_key("Md/Ro") == "MD-RO"
    assert corridor_key("Transfer") == "UA-MD-RO"
    assert corridor_key("Ro Price") is None
    assert corridor_key("Capacity won") is None


def test_position_sheet_becomes_bids_and_prices():
    intake = read_excel_ops(_operator_sheet(), day="2026-09-14")
    assert intake.warnings == []
    cap = intake.bids["capacity"]
    assert cap["RO-UA"][0] == 15.0 and cap["UA-MD-RO"][23] == 10.0
    assert cap["UA-MD"][5] == 0.0  # recorded as "not held", dropped on write
    assert intake.bids["cbc"]["RO-UA"][0] == 0.23 and intake.bids["cbc"]["RO-UA"][10] == 0.0
    assert intake.bids["limits"]["RO-UA"][2] == 58.0
    assert "filled" not in intake.bids
    # Ro/Ua Price columns are day prices, not position data.
    assert intake.prices_override["RO"][0] == 101.0
    assert intake.prices_override["UA"][23] == 95.5
    assert intake.availability == {}
    summary = intake.position_summary()
    assert summary[0].startswith("RO-UA 15 MW × 24h · CBC 0.00–0.23")
    assert summary[1].startswith("UA-MD-RO 10 MW × 24h")


def test_results_sheet_carries_fills_and_reported_pnl():
    intake = read_excel_ops(_operator_sheet(with_results=True), day="2026-09-14")
    assert intake.bids["filled"]["RO-UA"] == {0: 15.0, 1: 15.0, 2: 15.0} | {
        h: 0.0 for h in range(3, 24)
    }
    assert intake.bids["realized"]["RO-UA"][0] == 120.5
    assert intake.bids["realized"]["UA-MD-RO"][7] == -3.0
    assert any("prins 3h" in s and "rezultat raportat €362" in s for s in intake.position_summary())


def test_write_bids_merges_morning_position_with_evening_results(tmp_path: Path):
    path = tmp_path / "bids_2026-09-14.csv"
    morning = read_excel_ops(_operator_sheet(), day="2026-09-14")
    out = write_bids_csv(path, morning.bids, "2026-09-14")
    assert out["corridors"] == ["RO-UA", "UA-MD-RO"] and out["rows"] == 48
    won = load_bids_csv(path)
    assert won["capacity"]["RO-UA"][0] == 15.0 and won["limits"]["RO-UA"][0] == 56.0
    assert "UA-MD" not in won["capacity"]
    assert won["filled"] == {} and won["realized"] == {}
    # Re-posting the same sheet changes nothing.
    assert write_bids_csv(path, morning.bids, "2026-09-14")["changed"] == 0

    evening = read_excel_ops(_operator_sheet(with_results=True), day="2026-09-14")
    out = write_bids_csv(path, evening.bids, "2026-09-14")
    assert out["rows"] == 48 and out["changed"] > 0
    won = load_bids_csv(path)
    assert won["capacity"]["RO-UA"][5] == 15.0  # morning columns survive
    assert won["filled"]["RO-UA"][0] == 15.0 and won["filled"]["RO-UA"][5] == 0.0
    assert won["realized"]["RO-UA"][2] == 120.5
    header = path.read_text().splitlines()[0]
    assert header.startswith("delivery_day,corridor,hour_cet,capacity_mw,cbc_price_eur_mwh")
    assert "filled_mw,realized_eur,note" in header


def test_csv_export_of_the_same_sheet_is_understood():
    text = (
        b"Intervals CET;Capacity won;;CBC Price;;Bid Limit Price;;Ro Price\n"
        b";Ro - Ua;Transfer;Ro - Ua;Transfer;Ro - Ua;Transfer;\n"
        b"1;15;0;0,23;0;66;0;101,5\n"
        b"2;15;0;0,23;0;55;0;99\n"
        b"Total;30\n"
    )
    intake = read_csv_ops(text, day="2026-09-14")
    assert intake.bids["capacity"]["RO-UA"] == {0: 15.0, 1: 15.0}
    assert intake.bids["cbc"]["RO-UA"][0] == 0.23
    assert intake.bids["limits"]["RO-UA"][0] == 66.0
    assert intake.prices_override["RO"] == {0: 101.5, 1: 99.0}
    assert intake.warnings == []


def _summary_sheet() -> bytes:
    """The operator's monthly Import/Export sheet, as posted on 13.09 (values from it)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Export"
    ws.append(["", "Import / Export"])
    ws.append(["Septembrie 2026", "", "", "AZI Import/Export", "", "TOTAL Import/Export", ""])
    ws.append(["", "", "", "MW", "Profit", "MW", "Profit"])
    ws.append(["Ro - Ua", "", "", 0, "-7,35 €", 15, "-543,22 €"])
    ws.append(["Ua - Ro", 420, "Ua - Ro", 0, "0,00 €", 302, "5.707,03 €"])
    ws.append(["", "", "Ro - Md", 0, "0,00 €", 118, "-2.988,07 €"])
    ws.append(["Ua - Md", 2588, "Ua - Md", 0, "0,00 €", 812, "13.604,27 €"])
    ws.append(["", "", "Md - Ro", 15, "391,58 €", 1776, "28.053,15 €"])
    ws.append(["Md - Ua", 0, "Md - Ua", 0, "0,00 €", 0, "0,00 €"])
    ws.append(["", "", "Ro - Md", 0, "0,00 €", 0, "0,00 €"])
    ws.append(["Ro - Md", "", "", 0, "0,00 €", 0, "0,00 €"])
    ws.append(["Md - Ro", "", "", 0, "0,00 €", 320, "8.752,82 €"])
    ws.append(["TOTAL", "", "", "Astazi", "384,23 €", "Profit", "52.585,98 €"])
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_summary_sheet_gives_day_results_per_corridor(tmp_path: Path):
    from energy_trading.ops_intake import DAY_TOTAL

    intake = read_excel_ops(_summary_sheet(), day="2026-09-14")
    assert intake.warnings == []
    realized = intake.bids["realized"]
    assert realized["RO-UA"] == {DAY_TOTAL: -7.35}
    assert realized["UA-MD-RO"] == {DAY_TOTAL: 391.58}  # group Ua-Md + leg Md-Ro = transit
    assert realized["UA-RO"] == {DAY_TOTAL: 0.0} and realized["UA-RO-MD"] == {DAY_TOTAL: 0.0}
    assert realized["MD-RO"] == {DAY_TOTAL: 0.0}
    assert round(sum(v[DAY_TOTAL] for v in realized.values()), 2) == 384.23
    assert set(intake.bids) == {"realized"}

    path = tmp_path / "bids_2026-09-14.csv"
    write_bids_csv(path, intake.bids, "2026-09-14")
    lines = path.read_text().splitlines()
    assert any(line.startswith("2026-09-14,RO-UA,day,,,,,-7.35") for line in lines)
    won = load_bids_csv(path)
    assert won["realized"]["UA-MD-RO"] == {DAY_TOTAL: 391.58}
    assert won["capacity"] == {}  # a day-total row is not a held hour

    prices = {"RO": {h: 200.0 for h in range(24)}, "UA": {h: 100.0 for h in range(24)}}
    report = day_pnl("2026-09-14", prices, won, [], home="RO")
    assert report["corridors"] == []  # nothing held, yet the booked results are kept
    assert report["net"]["reported"] == 384.23
    assert report["reported_unmatched"]["UA-MD-RO"] == 391.58


def test_reported_pnl_shows_calibration_gap():
    prices = {"RO": {h: 100.0 for h in range(24)}, "UA": {h: 130.0 for h in range(24)}}
    won = {
        "capacity": {"RO-UA": {0: 15.0, 1: 15.0}},
        "cbc": {"RO-UA": {0: 0.0, 1: 0.0}},
        "limits": {},
        "filled": {"RO-UA": {0: 15.0, 1: 0.0}},
        "realized": {"RO-UA": {0: 300.0, 1: 0.0}},
    }
    report = day_pnl("2026-09-14", prices, won, [], home="RO")
    net = report["net"]
    assert net["operator"] is not None and net["reported"] == 300.0
    assert net["reported_gap"] == round(300.0 - net["operator"], 2)
    text = format_pnl(report)
    assert "Raportat de tine: €300" in text and "diferență față de calculul meu" in text


def test_telegram_document_writes_position_file_and_watch_reruns_pnl(tmp_path, monkeypatch):
    import shutil

    from energy_trading import scheduler as sched
    from energy_trading.agent import AgentConfig, CrossBorderAgent
    from energy_trading.config import Settings
    from energy_trading.scheduler import JobRunner
    from energy_trading.store import StateStore
    from energy_trading.telegram_bot import TelegramIngestor

    data = tmp_path / "data"
    data.mkdir()
    for f in ("prices_2026-09-13.csv", "ntc_2026-09-13.csv"):
        shutil.copy(Path("data") / f, data / f)
    settings = Settings(
        data_dir=data,
        state_dir=tmp_path / "state",
        telegram_bot_token="",
        telegram_chat_id="",
        fetch_minutes=0,
        pnl_hour=23,
    )
    store = StateStore(settings.state_dir)
    runner = JobRunner(CrossBorderAgent(config=AgentConfig()), settings, store)
    monkeypatch.setattr(sched, "FETCHERS", {})
    sent: list[str] = []
    runner.notifier.send = lambda text: sent.append(text) or True
    runner.run("watch")  # baseline

    ingestor = TelegramIngestor(settings, store)
    update = {
        "message": {
            "chat": {"id": "1"},
            "document": {"file_id": "x", "file_name": "operatiuni 13.09.xlsx"},
            "caption": "operatiuni 13.09.2026 realizate",
        }
    }
    sheet = _operator_sheet(with_results=True)
    res = ingestor.handle_update(update, download=lambda _fid: sheet)
    assert res.accepted and res.day == "2026-09-13"
    assert "📌 Poziție: RO-UA 15 MW × 24h" in res.reply
    assert (data / "bids_2026-09-13.csv").exists()

    out = runner.run("watch")
    assert out["result"]["pnl_reruns"] == ["2026-09-13"]
    assert any("P/L Digital Twin — 2026-09-13" in m and "Operațiuni reale" in m for m in sent)
    report = store.load("briefs/2026-09-13_pnl")
    assert report["net"]["reported"] is not None
    assert [c["corridor"] for c in report["corridors"]] == ["RO-UA", "UA-MD-RO"]
